import hashlib
import io
import math
import re
from collections import Counter
from dataclasses import dataclass
from typing import Iterable

import pymupdf
from PIL import Image, UnidentifiedImageError
from docx import Document
from django.utils import timezone
from django.conf import settings
from django.db.models import OuterRef, Subquery

from .models import EvidenceDocument, IngestionRun, KnowledgeChunk


TOKEN_PATTERN = re.compile(r"[a-zA-Z0-9₹]+(?:[.-][a-zA-Z0-9]+)*")
FIELD_QUERIES = {
    "maha_rera_id": "MahaRERA registration number project registration ID",
    "project_name": "project development name residential",
    "location": "address locality location city postal pin",
    "promoter_name": "developer promoter builder company",
    "configuration_types": "BHK configuration apartment residence",
    "carpet_area_sqft": "RERA carpet area square feet sqft",
    "room_dimensions": "bedroom living kitchen bathroom dimensions feet",
    "amenities": "amenities clubhouse gym pool garden parking security",
    "price": "price cost consideration INR crore lakh all inclusive",
    "possession_date": "possession completion handover date",
    "approvals": "approval registration MahaRERA permission certificate",
    "floor_plan": "floor plan unit plan layout room dimensions BHK carpet area",
}


@dataclass(frozen=True)
class SourcePage:
    page_number: int | None
    text: str
    metadata: dict


def sanitize_extracted_text(value):
    """Remove characters PostgreSQL text fields cannot store while retaining evidence text."""
    return str(value or "").replace("\x00", "")


class DocumentTextExtractor:
    """Extracts inert text only. It never follows links or executes embedded content."""

    def extract(self, document: EvidenceDocument, content: bytes, media_type=None) -> list[SourcePage]:
        media_type = media_type or document.detected_media_type or document.declared_media_type
        if media_type == "application/pdf":
            pymupdf.TOOLS.mupdf_display_errors(False)
            pymupdf.TOOLS.mupdf_display_warnings(False)
            try:
                with pymupdf.open(stream=content, filetype="pdf") as pdf:
                    if pdf.needs_pass:
                        raise ValueError("ENCRYPTED_PDF_UNSUPPORTED")
                    if pdf.page_count > settings.MAX_DOCUMENT_PAGES:
                        raise ValueError("DOCUMENT_PAGE_LIMIT_EXCEEDED")
                    return [
                        SourcePage(
                            page_number=index,
                            text=sanitize_extracted_text(page.get_text("text"))[:250_000],
                            metadata={"sourceType": "pdf_page", "imageCount": len(page.get_images(full=False))},
                        )
                        for index, page in enumerate(pdf, start=1)
                    ]
            finally:
                pymupdf.TOOLS.mupdf_display_errors(True)
                pymupdf.TOOLS.mupdf_display_warnings(True)
        if media_type in {"text/plain", "text/markdown", "text/csv", "application/json"}:
            text = sanitize_extracted_text(content.decode("utf-8-sig", errors="replace"))
            return [SourcePage(page_number=None, text=text[:1_000_000], metadata={"sourceType": media_type})]
        if media_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
            document = Document(io.BytesIO(content))
            blocks = [paragraph.text for paragraph in document.paragraphs]
            for table in document.tables:
                blocks.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
            return [SourcePage(page_number=None, text=sanitize_extracted_text("\n".join(blocks))[:1_000_000], metadata={"sourceType": "docx"})]
        if media_type in {"image/png", "image/jpeg"}:
            Image.MAX_IMAGE_PIXELS = 40_000_000
            try:
                with Image.open(io.BytesIO(content)) as image:
                    image.verify()
                with Image.open(io.BytesIO(content)) as image:
                    width, height = image.size
            except (UnidentifiedImageError, Image.DecompressionBombError) as exc:
                raise ValueError("INVALID_OR_UNSAFE_IMAGE") from exc
            return [SourcePage(
                page_number=1,
                text="",
                metadata={"sourceType": "image", "requiresVision": True, "width": width, "height": height},
            )]
        raise ValueError("UNSUPPORTED_MEDIA_TYPE")


def ingest_document(document, content, pipeline_version="rag-ingestion-1.0", pages=None, defer_completion=False):
    run = IngestionRun.objects.create(
        project=document.project,
        document=document,
        status="extracting",
        pipeline_version=pipeline_version,
        started_at=timezone.now(),
    )
    try:
        pages = DocumentTextExtractor().extract(document, content) if pages is None else pages
        chunks = chunk_pages(pages)
        KnowledgeChunk.objects.bulk_create([
            KnowledgeChunk(
                project=document.project,
                document=document,
                ingestion_run=run,
                ordinal=index,
                page_number=chunk.page_number,
                content=chunk.text,
                content_hash=hashlib.sha256(chunk.text.encode()).hexdigest(),
                token_count=len(tokenize(chunk.text)),
                embedding=sparse_embedding(chunk.text),
                metadata=chunk.metadata,
            )
            for index, chunk in enumerate(chunks, start=1)
        ])
        run.status = "extracting" if defer_completion else "completed"
        run.completed_at = timezone.now()
        run.summary = {
            "pageCount": len(pages),
            "chunkCount": len(chunks),
            "textCharacters": sum(len(page.text) for page in pages),
            "visionPages": [page.page_number for page in pages if page.metadata.get("requiresVision") or page.metadata.get("imageCount", 0)],
        }
        run.save(update_fields=["status", "completed_at", "summary", "updated_at"])
        return run
    except Exception as exc:
        run.status = "failed"
        run.error_code = str(exc)[:64] or exc.__class__.__name__
        run.save(update_fields=["status", "error_code", "updated_at"])
        raise


def chunk_pages(pages: Iterable[SourcePage], maximum_words=280, overlap_words=40):
    chunks = []
    stride = maximum_words - overlap_words
    for page in pages:
        safe_text = sanitize_extracted_text(page.text)
        words = safe_text.split()
        if not words:
            chunks.append(SourcePage(page.page_number, safe_text, page.metadata))
            continue
        for offset in range(0, len(words), stride):
            text = " ".join(words[offset:offset + maximum_words]).strip()
            if text:
                chunks.append(SourcePage(page.page_number, text, page.metadata | {"wordOffset": offset}))
            if offset + maximum_words >= len(words):
                break
    return chunks


def retrieve(project, query, limit=6, ingestion_run=None):
    query_embedding = sparse_embedding(query)
    chunks = KnowledgeChunk.objects.filter(project=project)
    if ingestion_run is not None:
        chunks = chunks.filter(ingestion_run=ingestion_run)
    else:
        latest_successful_run = IngestionRun.objects.filter(
            document_id=OuterRef("document_id"),
            status="completed",
            chunks__isnull=False,
        ).order_by("-completed_at", "-created_at").values("id")[:1]
        chunks = chunks.filter(ingestion_run_id=Subquery(latest_successful_run))
    scored = []
    for chunk in chunks.iterator():
        score = cosine_similarity(query_embedding, chunk.embedding)
        if score > 0:
            scored.append((score, chunk))
    scored.sort(key=lambda item: (-item[0], item[1].ordinal))
    return [{"score": score, "chunk": chunk} for score, chunk in scored[:limit]]


def retrieve_field_context(project, ingestion_run):
    return {
        field_code: retrieve(project, query, limit=4, ingestion_run=ingestion_run)
        for field_code, query in FIELD_QUERIES.items()
    }


def tokenize(text):
    return [token.casefold() for token in TOKEN_PATTERN.findall(text)]


def sparse_embedding(text, dimensions=384):
    counts = Counter(tokenize(text))
    vector = {}
    for token, count in counts.items():
        index = int(hashlib.sha256(token.encode()).hexdigest()[:8], 16) % dimensions
        vector[str(index)] = vector.get(str(index), 0.0) + 1.0 + math.log(count)
    norm = math.sqrt(sum(value * value for value in vector.values())) or 1.0
    return {key: round(value / norm, 8) for key, value in vector.items()}


def cosine_similarity(left, right):
    if not left or not right:
        return 0.0
    return sum(value * right.get(key, 0) for key, value in left.items())
